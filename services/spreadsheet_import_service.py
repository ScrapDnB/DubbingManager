"""Spreadsheet readers and row conversion, independent of Qt and project state."""

from __future__ import annotations

import codecs
import csv
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
import io
import math
from pathlib import Path
import re
from typing import Any

from services.docx_import_service import COLUMN_TYPES, DocxImportService
from utils.helpers import format_seconds_to_full_tc


SPREADSHEET_EXTENSIONS = {".xls", ".xlsx", ".csv"}


@dataclass
class SpreadsheetSheet:
    name: str
    rows: list[list[Any]]
    hidden: bool = False

    @property
    def column_count(self) -> int:
        return max((len(row) for row in self.rows), default=0)


class SpreadsheetImportService:
    """Keep physical rows/columns so saved import settings remain reproducible."""

    def __init__(self) -> None:
        self.encoding = ""
        self.delimiter = ""

    def load(self, path: str, encoding: str = "auto", delimiter: str = "auto") -> list[SpreadsheetSheet]:
        suffix = Path(path).suffix.lower()
        self.encoding = self.delimiter = ""
        if suffix == ".csv":
            return self._load_csv(path, encoding, delimiter)
        if suffix == ".xlsx":
            return self._load_xlsx(path)
        if suffix == ".xls":
            return self._load_xls(path)
        raise ValueError("Выберите файл XLS, XLSX или CSV")

    def _load_csv(self, path: str, encoding: str, delimiter: str) -> list[SpreadsheetSheet]:
        raw = Path(path).read_bytes()
        if encoding == "auto":
            if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
                encoding = "utf-16"
            else:
                try:
                    raw.decode("utf-8-sig")
                    encoding = "utf-8-sig"
                except UnicodeDecodeError:
                    encoding = "cp1251"
        content = raw.decode(encoding).lstrip("\ufeff")
        if "\x00" in content:
            raise ValueError("Не удалось прочитать CSV: проверьте кодировку")
        # Excel's optional delimiter directive is metadata, not a data row.
        first, _, rest = content.partition("\n")
        directive = re.fullmatch(r"sep=(.)", first.rstrip("\r"), re.IGNORECASE)
        if directive:
            content = rest
        if delimiter == "auto":
            if directive:
                delimiter = directive.group(1)
            else:
                try:
                    delimiter = csv.Sniffer().sniff(content[:65536], delimiters=",;\t|").delimiter
                except csv.Error:
                    delimiter = ","
        if len(delimiter) != 1:
            raise ValueError("Разделитель CSV должен быть одним символом")
        rows = list(csv.reader(io.StringIO(content, newline=""), delimiter=delimiter, strict=True))
        self.encoding, self.delimiter = encoding, delimiter
        return [SpreadsheetSheet(Path(path).stem, self._trim_rows(rows))]

    def _load_xlsx(self, path: str) -> list[SpreadsheetSheet]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("Для импорта XLSX установите openpyxl") from exc
        # data_only uses the results saved by Excel; formulas are never executed.
        workbook = load_workbook(path, data_only=True, keep_links=False)
        try:
            sheets = []
            for sheet in workbook.worksheets:
                rows = [[None if cell.data_type == "e" else cell.value for cell in row]
                        for row in sheet.iter_rows()]
                for merged in sheet.merged_cells.ranges:
                    value = rows[merged.min_row - 1][merged.min_col - 1]
                    for row in range(merged.min_row - 1, merged.max_row):
                        for column in range(merged.min_col - 1, merged.max_col):
                            rows[row][column] = value
                sheets.append(SpreadsheetSheet(sheet.title, self._trim_rows(rows), sheet.sheet_state != "visible"))
            return sheets
        finally:
            workbook.close()

    def _load_xls(self, path: str) -> list[SpreadsheetSheet]:
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("Для импорта XLS установите xlrd") from exc
        workbook = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
        try:
            sheets = []
            for sheet in workbook.sheets():
                rows = []
                for row_index in range(sheet.nrows):
                    row = []
                    for cell in sheet.row(row_index):
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            # Excel durations, unlike dates, may extend beyond 24h.
                            format_key = workbook.xf_list[cell.xf_index].format_key
                            format_string = workbook.format_map[format_key].format_str.casefold()
                            if 0 <= value < 1 or "[h]" in format_string or "[m]" in format_string or "[s]" in format_string:
                                value = timedelta(days=value)
                            else:
                                value = xlrd.xldate_as_datetime(value, workbook.datemode)
                        elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                            value = None
                        row.append(value)
                    rows.append(row)
                for rlo, rhi, clo, chi in sheet.merged_cells:
                    for row in range(rlo, rhi):
                        for column in range(clo, chi):
                            rows[row][column] = rows[rlo][clo]
                sheets.append(SpreadsheetSheet(sheet.name, self._trim_rows(rows), bool(sheet.visibility)))
            return sheets
        finally:
            workbook.release_resources()

    @staticmethod
    def _trim_rows(rows: list[list[Any]]) -> list[list[Any]]:
        # Never remove leading/interior blank rows or columns.
        while rows and not any(value is not None and str(value).strip() for value in rows[-1]):
            rows.pop()
        for row in rows:
            while row and (row[-1] is None or row[-1] == ""):
                row.pop()
        return rows

    @staticmethod
    def cell_text(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        if isinstance(value, (datetime, time)):
            return value.strftime("%H:%M:%S") + (f".{value.microsecond:06d}" if value.microsecond else "")
        return str(value).strip()

    def detect(self, sheet: SpreadsheetSheet) -> dict:
        helper = DocxImportService()
        rows = [[self.cell_text(value) for value in row] for row in sheet.rows]
        candidates = [(i, self._header_mapping(row, helper)) for i, row in enumerate(rows[:20])]
        header_index, header = max(candidates, key=lambda item: len(item[1]), default=(None, {}))
        text_index = header.get("text")
        exact_text_header = (text_index is not None and header_index is not None
                             and rows[header_index][text_index].casefold() in helper.detection_config["aliases"]["text"])
        if text_index is None or (len(header) < 2 and not exact_text_header):
            header_index = None
        detected = helper.detect_columns(rows)
        # Header fields must not inherit fallback columns that belong to other fields.
        if header_index is not None:
            detected = header
        mapping = {key: detected.get(key) if isinstance(detected.get(key), int)
                   and 0 <= detected[key] < sheet.column_count else None for key in COLUMN_TYPES}
        if header_index is None and sheet.column_count < 3:
            mapping = {key: None for key in COLUMN_TYPES}
            mapping["text"] = sheet.column_count - 1 if sheet.column_count else None
            if sheet.column_count == 2:
                mapping["character"] = 0
        return {"mapping": mapping, "first_row": (header_index + 2) if header_index is not None else 1,
                "ignore_header": header_index is not None,
                "time_unit": "seconds", "fps": 25.0, "default_duration": 1.0,
                "time_separators": ["-", "–", "—", "|", "-->"]}

    @staticmethod
    def _header_mapping(row: list[str], helper: DocxImportService) -> dict:
        mapping, used = {}, set()
        for field in COLUMN_TYPES:
            for i, cell in enumerate(row):
                normalized = " ".join(cell.casefold().split())
                if i not in used and any(
                    normalized == alias if len(alias) <= 3 else alias in normalized
                    for alias in helper.detection_config["aliases"].get(field, [])
                ):
                    mapping[field] = i
                    used.add(i)
                    break
        return mapping

    def parse_time(self, value: Any, config: dict) -> float | None:
        if value is None or value == "":
            return None
        if isinstance(value, timedelta):
            seconds = value.total_seconds()
        elif isinstance(value, (datetime, time)):
            seconds = value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
        elif isinstance(value, date):
            return None
        else:
            raw = str(value).strip().replace(",", ".")
            if re.fullmatch(r"\d+(?:\.\d+)?", raw):
                seconds = float(raw) * (86400 if config.get("time_unit") == "excel" else 1)
            else:
                frames = re.fullmatch(r"(\d+):(\d{2}):(\d{2}):(\d{2})", raw)
                clock = re.fullmatch(r"(?:(\d+):)?(\d{1,2}):(\d{2})(?:\.(\d+))?", raw)
                if frames:
                    h, m, s, f = map(int, frames.groups())
                    fps = float(config.get("fps", 25))
                    if fps <= 0 or m >= 60 or s >= 60 or f >= fps:
                        return None
                    seconds = h * 3600 + m * 60 + s + f / fps
                elif clock:
                    h, m, s, fraction = clock.groups()
                    if int(s) >= 60 or (h is not None and int(m) >= 60):
                        return None
                    seconds = int(h or 0) * 3600 + int(m) * 60 + int(s) + float("0." + (fraction or "0"))
                else:
                    return None
        return seconds if math.isfinite(seconds) and seconds >= 0 else None

    def parse_sheet(self, sheet: SpreadsheetSheet, config: dict, limit: int | None = None) -> tuple[list[dict], list[dict]]:
        mapping = config.get("mapping", {})
        if not DocxImportService.mapping_usable(mapping, sheet.column_count):
            raise ValueError(f"Лист «{sheet.name}»: укажите корректные колонки, включая текст")
        first = int(config.get("first_row", 1))
        if not 1 <= first <= len(sheet.rows):
            raise ValueError(f"Лист «{sheet.name}»: первая строка данных вне листа")
        duration = float(config.get("default_duration", 1))
        if not math.isfinite(duration) or duration <= 0:
            raise ValueError("Длительность реплики должна быть больше нуля")
        lines, preview = [], []
        rows = sheet.rows[first - 1:]
        if limit is not None:
            rows = rows[:limit]
        for number, row in enumerate(rows, first):
            def cell(field):
                index = mapping.get(field)
                return row[index] if index is not None and index < len(row) else None
            text = self.cell_text(cell("text"))
            character = self.cell_text(cell("character"))
            start_raw, end_raw = cell("time_start"), cell("time_end")
            combined = self.cell_text(cell("time_split"))
            if combined:
                for separator in sorted(config.get("time_separators", ["-", "–", "—", "|", "-->"]), key=len, reverse=True):
                    if separator and separator in combined:
                        start_raw, end_raw = combined.split(separator, 1)
                        break
                else:
                    start_raw = combined
            start, end = self.parse_time(start_raw, config), self.parse_time(end_raw, config)
            invalid = ((self.cell_text(start_raw) and start is None)
                       or (self.cell_text(end_raw) and end is None))
            start = 0.0 if start is None else start
            end = start + duration if end is None else end
            invalid = bool(invalid or end < start)
            status = "Нет текста" if not text else "Ошибка тайминга" if invalid else "Готово"
            item = {"row": number, "character": character, "text": text,
                    "timing": f"{format_seconds_to_full_tc(start)} – {format_seconds_to_full_tc(end)}", "status": status}
            preview.append(item)
            if text and not invalid:
                lines.append({"s": start, "e": end, "char": character, "text": text,
                              "s_raw": self.cell_text(start_raw), "e_raw": self.cell_text(end_raw)})
        return lines, preview

    def parse_source(self, path: str, config: dict) -> list[dict]:
        sheets = self.load(path, config.get("encoding", "auto"), config.get("delimiter", "auto"))
        name = config.get("sheet_name")
        if name is None:
            if len(sheets) != 1:
                raise ValueError("Укажите лист для восстановления текста")
            sheet = sheets[0]
        else:
            sheet = next((sheet for sheet in sheets if sheet.name == name), None)
            if sheet is None:
                raise ValueError(f"В файле больше нет листа «{name}»")
        options = config if config.get("mapping") else {**self.detect(sheet), **config}
        lines, preview = self.parse_sheet(sheet, options)
        errors = [item for item in preview if item["status"] == "Ошибка тайминга"]
        if errors:
            raise ValueError(f"Лист «{sheet.name}», строка {errors[0]['row']}: ошибка тайминга")
        return lines
