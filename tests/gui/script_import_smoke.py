"""Isolated DOCX/spreadsheet window smoke and visual checks on macOS.

Run .venv/bin/python tests/gui/script_import_smoke.py [--fluent] [--dark].
Fluent on macOS checks layout only; it is not Windows runtime validation.
"""
import json
from pathlib import Path
import sys
import tempfile

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from PySide6.QtCore import QCoreApplication, QEvent, QObject, QPoint, QSettings, QTimer, QUrl, Qt, qInstallMessageHandler
from PySide6.QtGui import QColor, QGuiApplication
from PySide6.QtQml import QQmlComponent, QQmlEngine
from PySide6.QtQuick import QQuickItem, QQuickWindow
from PySide6.QtQuickControls2 import QQuickStyle
from PySide6.QtTest import QTest
from shiboken6 import getCppPointer, wrapInstance
from docx import Document
from openpyxl import Workbook
import services.global_settings_service as gs
from ui.macos_integration import MacOSIntegration
from ui.qml_backend.app_bridge import AppBridge
from ui.qml_backend.features.ui_state_bridge import UiStateBridge

root = Path(__file__).resolve().parents[2]
folder = tempfile.TemporaryDirectory(prefix='dm-import-smoke-')
tmp = Path(folder.name)
gs.SETTINGS_FILE = tmp / 'settings.json'
fluent = '--fluent' in sys.argv
dark = '--dark' in sys.argv
QQuickStyle.setStyle('FluentWinUI3' if fluent else 'macOS')
app = QGuiApplication(sys.argv)
ui = UiStateBridge(QSettings(str(tmp / 'ui.ini'), QSettings.IniFormat))
ui.setStringValue('main.theme', 'dark' if dark else 'light')
bridge = AppBridge(ui_state=ui)
platform = MacOSIntegration()
engine = QQmlEngine()
engine.rootContext().setContextProperty('platformIntegration', platform)
owner = QQuickWindow()
owner.resize(1400, 1040)
errors, failures = [], []

def message(kind, context, text):
    if 'file:' in text or 'ReferenceError' in text or 'TypeError' in text or 'Binding loop' in text:
        errors.append(text)
qInstallMessageHandler(message)
output = root / 'output/import-smoke'
output.mkdir(parents=True, exist_ok=True)
prefix = ('fluent' if fluent else 'macos') + ('-dark' if dark else '-light')
colors = {'softBorder': '#565656' if dark else '#d7dce2', 'softHeader': '#303236' if dark else '#f0f3f7',
          'softRow': '#25272b' if dark else '#ffffff', 'softAltRow': '#2a2c30' if dark else '#f6f8fb',
          'softMuted': '#afb5bf' if dark else '#697586'}
dialogs = []
qml_objects = []

def check(value, text):
    if not value:
        failures.append(text)
    print(('PASS ' if value else 'FAIL ') + text, flush=True)

def invoke(obj, expression):
    engine.globalObject().setProperty('testDialog', engine.newQObject(obj))
    result = engine.evaluate('testDialog.' + expression)
    check(not result.isError(), expression + (': ' + result.toString() if result.isError() else ''))
    return result

def create(name):
    component = QQmlComponent(engine, QUrl.fromLocalFile(str(root / 'qml/components' / name)))
    props = {'appBridge': bridge, 'ownerWindow': owner, **{k: QColor(v) for k,v in colors.items()}}
    window = component.createWithInitialProperties(props)
    check(window is not None, 'create ' + name + ': ' + '\n'.join(e.toString() for e in component.errors()))
    if window is None:
        app.exit(1)
        return None
    qml_objects.extend([component, window])
    window = wrapInstance(getCppPointer(window)[0], QQuickWindow)
    # Opaque capture background; native owner styling is tested by the main app smoke.
    window.setColor(QColor('#25272b' if dark else '#ffffff'))
    dialogs.append(window)
    return window

def check_geometry(window):
    for name in ['importColumn_text', 'importColumn_character', 'importEpisodeName', 'importConfirmButton', 'importPreview']:
        pending = [window.contentItem()]
        item = None
        while pending:
            candidate = pending.pop()
            pending.extend(candidate.childItems())
            if candidate.objectName() == name:
                item = candidate
                break
        check(item is not None, name + ' exists')
        if item is None:
            continue
        item = wrapInstance(getCppPointer(item)[0], QQuickItem)
        point = item.mapToScene(item.boundingRect().topLeft())
        check(point.x() >= 0 and point.x() + item.width() <= window.width() + 1
              and (name.startswith('importColumn_') or (point.y() >= 0 and point.y() + item.height() <= window.height() + 1)),
              name + ' stays inside window')

book = Workbook()
sheet = book.active
sheet.title = '01 · Северный маяк'
sheet.append(['Северный маяк · серия 01'])
sheet.append(['Персонаж','Начало','Конец','Текст'])
for values in [('Алиса', '00:00:02', '00:00:05', 'Ты слышишь? Маяк снова работает.'),
               ('Марк', '00:00:06', '00:00:10', 'Я думал, его отключили много лет назад.'),
               ('Рассказчик', '00:00:11', '00:00:17', 'Свет маяка медленно скользнул по воде.\nУ самого берега показалась лодка.')]:
    sheet.append(values)
sheet = book.create_sheet('02 · Возвращение')
sheet.append(['Текст','Персонаж','Тайминг'])
sheet.append(['Это началось вчера вечером.', 'Алиса', '00:00:01 – 00:00:04'])
book.create_sheet('Примечания')
xlsx = tmp / 'Северный маяк.xlsx'
book.save(xlsx)
doc = Document()
for title in ['Ты слышишь? Маяк снова работает.', 'Посмотрите на воду. Там кто-то есть.']:
    table = doc.add_table(rows=4, cols=3)
    for r, values in enumerate([['Персонаж','Тайминг','Текст'], ['Алиса','00:00:02 – 00:00:05',title],
                               ['Марк','00:00:06 – 00:00:10','Я думал, его отключили много лет назад.'],
                               ['Рассказчик','00:00:11 – 00:00:17','Свет маяка медленно скользнул по воде.']]):
        for c,v in enumerate(values):
            table.cell(r,c).text = v
docx = tmp / 'Северный маяк 01.docx'
doc.save(docx)

spreadsheet = create('SpreadsheetImportDialog.qml')
if spreadsheet:
    invoke(spreadsheet, 'openForFile(' + json.dumps(str(xlsx), ensure_ascii=False) + ')')

def capture_spreadsheet():
    check_geometry(spreadsheet)
    spreadsheet.grabWindow().save(str(output / (prefix + '-xlsx.png')))
    check(bridge.spreadsheetImport.canImport, 'selected sheets ready')
    # Activate the second real sheet card with a pointer event.
    cards = spreadsheet.findChild(QObject, 'importSheets')
    cards = wrapInstance(getCppPointer(cards)[0], QQuickItem)
    point = cards.mapToScene(QPoint(340, 22))
    QTest.mouseClick(spreadsheet, Qt.LeftButton, Qt.NoModifier, QPoint(round(point.x()),round(point.y())))
    check(bridge.spreadsheetImport.currentSheetIndex == 1, 'sheet card navigates')
    check(bridge.spreadsheetImport.mapping['text'] == 0, 'sheet-specific columns restored')
    QTimer.singleShot(250, capture_second)

def capture_second():
    name = spreadsheet.findChild(QObject, 'importEpisodeName')
    check(name.property('text') == '02 · Возвращение', 'series field follows sheet selection')
    # Click the header checkbox twice and the shared-column action.
    checkbox = wrapInstance(getCppPointer(spreadsheet.findChild(QObject, 'importIgnoreHeader'))[0], QQuickItem)
    point = checkbox.mapToScene(checkbox.boundingRect().center())
    QTest.mouseClick(spreadsheet, Qt.LeftButton, Qt.NoModifier, QPoint(round(point.x()), round(point.y())))
    check(not bridge.spreadsheetImport.ignoreHeader and bridge.spreadsheetImport.firstRow == 1,
          'header checkbox includes the header when unchecked')
    QTest.mouseClick(spreadsheet, Qt.LeftButton, Qt.NoModifier, QPoint(round(point.x()), round(point.y())))
    check(bridge.spreadsheetImport.ignoreHeader and bridge.spreadsheetImport.firstRow == 2,
          'header checkbox excludes the header when checked')
    columns = bridge.spreadsheetImport.mapping
    button = wrapInstance(getCppPointer(spreadsheet.findChild(QObject, 'importApplyColumnsToAll'))[0], QQuickItem)
    point = button.mapToScene(button.boundingRect().center())
    QTest.mouseClick(spreadsheet, Qt.LeftButton, Qt.NoModifier, QPoint(round(point.x()), round(point.y())))
    bridge.spreadsheetImport.setSheet(0)
    check(bridge.spreadsheetImport.mapping == columns, 'apply-all button copies columns to another series')
    check(bridge.spreadsheetImport.tableStartRow == 2, 'apply-all retains the other series table start')
    check(bridge.spreadsheetImport.ignoreHeader, 'apply-all copies the header checkbox')
    bridge.spreadsheetImport.setSheet(1)
    spreadsheet.resize(980, 680)
    QTimer.singleShot(250, capture_compact)

def capture_compact():
    check_geometry(spreadsheet)
    spreadsheet.grabWindow().save(str(output / (prefix + '-xlsx-compact.png')))
    confirm = wrapInstance(getCppPointer(spreadsheet.findChild(QObject,'importConfirmButton'))[0], QQuickItem)
    point = confirm.mapToScene(confirm.boundingRect().center())
    QTest.mouseClick(spreadsheet, Qt.LeftButton, Qt.NoModifier, QPoint(round(point.x()),round(point.y())))
    check(not spreadsheet.isVisible(), 'import closes window')
    check(len(bridge._session.data['episodes']) == 2, 'button imports selected sheets')
    bridge.project.undo()
    check(not bridge._session.data['episodes'], 'one undo removes whole workbook')
    document = create('DocxImportDialog.qml')
    invoke(document, 'openForFile(' + json.dumps(str(docx), ensure_ascii=False) + ')')
    QTimer.singleShot(250, lambda: capture_docx(document))

def capture_docx(document):
    check_geometry(document)
    document.grabWindow().save(str(output / (prefix + '-docx.png')))
    # Locate the common raw-preview tab through the visual tree.
    pending = [document.contentItem()]
    while pending:
        item = pending.pop()
        pending.extend(item.childItems())
        if item.property('text') == 'Исходные колонки':
            point = item.mapToScene(item.boundingRect().center())
            QTest.mouseClick(document, Qt.LeftButton, Qt.NoModifier, QPoint(round(point.x()),round(point.y())))
            break
    QTimer.singleShot(250, lambda: capture_raw(document))

def capture_raw(document):
    document.grabWindow().save(str(output / (prefix + '-docx-raw.png')))
    document.close()
    check(not document.isVisible(), 'window close dismisses DOCX')
    check(not bridge._session.data['episodes'], 'closing preview does not mutate project')
    csv = tmp / 'Серия 03.csv'
    csv.write_text('Персонаж;Текст\nАлиса;Мы вернулись!\n', encoding='cp1251')
    spreadsheet.resize(1220, 860)
    invoke(spreadsheet, 'openForFile(' + json.dumps(str(csv), ensure_ascii=False) + ')')
    QTimer.singleShot(250, finish)

def finish():
    check_geometry(spreadsheet)
    spreadsheet.grabWindow().save(str(output / (prefix + '-csv.png')))
    spreadsheet.close()
    check(not spreadsheet.isVisible(), 'window close dismisses CSV')
    check(not errors, 'no QML runtime errors: ' + repr(errors[:8]))
    (output / (prefix + '-checks.json')).write_text(json.dumps({'failures': failures, 'qmlErrors': errors},ensure_ascii=False,indent=2))
    app.exit(1 if failures or errors else 0)

QTimer.singleShot(450, capture_spreadsheet)
QTimer.singleShot(15000, lambda: app.exit(3))
code = app.exec()
for window in dialogs:
    window.deleteLater()
QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
qInstallMessageHandler(None)
bridge.audiobook.shutdown()
sys.exit(code)
