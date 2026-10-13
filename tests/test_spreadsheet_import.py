"""Real file import, per-sheet configuration, project persistence and undo."""
from copy import deepcopy
from datetime import time, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook
from PySide6.QtCore import QSettings

from services.spreadsheet_import_service import SpreadsheetImportService, SpreadsheetSheet
from ui.qml_backend.app_bridge import AppBridge
from ui.qml_backend.features.ui_state_bridge import UiStateBridge


@pytest.fixture
def bridge(tmp_path, monkeypatch):
    monkeypatch.setattr('services.global_settings_service.SETTINGS_FILE', tmp_path / 'settings.json')
    app = AppBridge(ui_state=UiStateBridge(QSettings(str(tmp_path / 'ui.ini'), QSettings.IniFormat)))
    app._project_service.set_backup_config({'enabled': False})
    yield app
    app.audiobook.shutdown()


def make_workbook(path):
    book = Workbook()
    sheet = book.active
    sheet.title = '01 · Начало'
    sheet.append(['Название серии'])
    sheet.append(['Персонаж', 'Начало', 'Конец', 'Текст'])
    sheet.append(['Герой', time(0, 0, 2), time(0, 0, 4), 'Первая реплика'])
    sheet.append([None, time(0, 0, 4), time(0, 0, 6), 'Вторая\nреплика'])
    sheet.merge_cells('A3:A4')
    second = book.create_sheet('02 · Продолжение')
    second.append(['Текст', 'Персонаж', 'Тайминг'])
    second.append(['Другой лист', 'Гость', '00:00:10 – 00:00:12'])
    hidden = book.create_sheet('Справочник')
    hidden.sheet_state = 'hidden'
    hidden.append(['Текст'])
    hidden.append(['Примечание'])
    book.create_sheet('Пустой')
    book.save(path)
    return path


@pytest.mark.parametrize('encoding,delimiter', [('utf-8-sig', ','), ('cp1251', ';'), ('utf-16', '\t')])
def test_csv_encodings_and_quoted_multiline(tmp_path, encoding, delimiter):
    import csv
    path = tmp_path / 'Episode_04.CSV'
    with path.open('w', encoding=encoding, newline='') as stream:
        writer = csv.writer(stream, delimiter=delimiter)
        writer.writerow(['Персонаж', 'Начало', 'Конец', 'Текст'])
        writer.writerow(['Герой', '1,5', '3,25', 'Привет, мир;\nещё строка'])
    service = SpreadsheetImportService()
    sheet, = service.load(str(path))
    config = service.detect(sheet)
    lines, preview = service.parse_sheet(sheet, config)
    assert service.delimiter == delimiter
    assert config['first_row'] == 2
    assert lines[0]['s'] == 1.5
    assert lines[0]['e'] == 3.25
    assert lines[0]['text'] == 'Привет, мир;\nещё строка'
    assert preview[0]['row'] == 2


def test_csv_sep_directive_and_physical_empty_rows(tmp_path):
    path = tmp_path / 'data.csv'
    path.write_text('sep=;\n\nЗаголовок файла\nПерсонаж;Текст\nГерой;Реплика\n\n', encoding='utf-8')
    service = SpreadsheetImportService()
    sheet, = service.load(str(path))
    assert len(sheet.rows) == 4
    config = service.detect(sheet)
    assert config['first_row'] == 4
    assert service.parse_sheet(sheet, config)[0][0]['text'] == 'Реплика'


def test_csv_wrong_encoding_is_reported(tmp_path):
    path = tmp_path / 'data.csv'
    path.write_text('Текст\nПривет', encoding='cp1251')
    with pytest.raises(UnicodeDecodeError):
        SpreadsheetImportService().load(str(path), 'utf-8')


def test_no_header_does_not_lose_first_line_and_skips_empty_text():
    service = SpreadsheetImportService()
    sheet = SpreadsheetSheet('data', [['Hero', '00:01 - 00:02', 'First line'],
                                     ['Guest', '00:03 - 00:04', 'Second line'], ['Note', '', '']])
    config = service.detect(sheet)
    assert config['first_row'] == 1
    lines, preview = service.parse_sheet(sheet, config)
    assert len(lines) == 2
    assert preview[-1]['status'] == 'Нет текста'


def test_xlsx_sheets_merged_names_and_excel_times(tmp_path):
    path = make_workbook(tmp_path / 'episodes.xlsx')
    service = SpreadsheetImportService()
    sheets = service.load(str(path))
    assert [s.hidden for s in sheets] == [False, False, True, False]
    assert sheets[-1].rows == []
    first = service.detect(sheets[0])
    assert first['first_row'] == 3
    assert first['mapping']['time_split'] is None
    lines, _ = service.parse_sheet(sheets[0], first)
    assert [l['char'] for l in lines] == ['Герой', 'Герой']
    assert [l['s'] for l in lines] == [2, 4]
    second = service.detect(sheets[1])
    assert second['mapping']['text'] == 0
    assert service.parse_sheet(sheets[1], second)[0][0]['s'] == 10


def test_real_xls_reader():
    service = SpreadsheetImportService()
    sheets = service.load(str(Path(__file__).parent / 'fixtures/spreadsheet_import.xls'))
    assert [s.name for s in sheets] == ['Серия 01', 'Серия 02', 'Скрытый', 'Пустой']
    assert sheets[2].hidden
    assert not sheets[3].rows
    lines, _ = service.parse_sheet(sheets[0], service.detect(sheets[0]))
    assert [l['char'] for l in lines] == ['Герой', 'Герой']
    assert lines[0]['s'] == pytest.approx(2)
    assert lines[1]['e'] == pytest.approx(4)


@pytest.mark.parametrize('value,options,expected', [
    ('00:01:02:12', {'fps': 24}, 62.5), ('01:02,500', {}, 62.5),
    (12.5, {}, 12.5), ('0.5', {'time_unit': 'excel'}, 43200),
    (time(1, 2, 3, 500000), {}, 3723.5), (timedelta(hours=27), {}, 97200),
    ('00:01:02:30', {'fps': 25}, None), ('00:61:00', {}, None),
    ('invalid', {}, None), (float('nan'), {}, None), ('-1', {}, None),
])
def test_time_formats(value, options, expected):
    assert SpreadsheetImportService().parse_time(value, options) == expected


def test_bad_timing_is_visible_and_not_silently_imported():
    service = SpreadsheetImportService()
    sheet = SpreadsheetSheet('data', [['Персонаж', 'Начало', 'Конец', 'Текст'],
                                     ['Герой', 'bad', 2, 'A'], ['Герой', 4, 2, 'B']])
    lines, preview = service.parse_sheet(sheet, service.detect(sheet))
    assert not lines
    assert [r['status'] for r in preview] == ['Ошибка тайминга', 'Ошибка тайминга']


def test_multiple_sheets_import_undo_redo_save_and_no_source(bridge, tmp_path):
    path = make_workbook(tmp_path / 'episodes.xlsx')
    importer = bridge.spreadsheetImport
    assert importer.load(path.as_uri())
    assert [s['selected'] for s in importer.sheets] == [True, True, False, False]
    importer.setEpisodeName('Первая')
    importer.setFirstRow(4)
    importer.setSheet(1)
    assert importer.mapping['text'] == 0
    importer.setEpisodeName('Вторая')
    importer.setSheet(0)
    assert importer.firstRow == 4
    assert importer.episodeName == 'Первая'
    assert importer.importSelected()
    data = bridge._session.data
    assert list(data['episodes']) == ['Первая', 'Вторая']
    assert bridge.project.currentEpisode == 'Первая'
    assert bridge._script_text_service.load_episode_lines(data, 'Первая')[0]['text'] == 'Вторая\nреплика'
    source = data['script_storage']['episodes']['Вторая']['source']
    assert source['line_mode'] == 'premerged'
    assert source['type'] == 'xlsx'
    assert source['import_config']['sheet_name'] == '02 · Продолжение'
    assert source['import_config']['mapping']['text'] == 0
    imported = deepcopy(data)
    bridge.project.undo()
    assert not data['episodes']
    bridge.project.redo()
    assert data == imported
    bridge.project.saveAs(str(tmp_path / 'project.dub'))
    saved = bridge.project.path
    path.unlink()
    bridge.project.open(saved)
    assert bridge._script_text_service.load_episode_lines(bridge._session.data, 'Вторая')[0]['text'] == 'Другой лист'


def test_sheet_deselection_and_name_collisions(bridge, tmp_path):
    path = make_workbook(tmp_path / 'episodes.xlsx')
    importer = bridge.spreadsheetImport
    assert importer.load(str(path))
    importer.setEpisodeName('Общая')
    importer.setSheet(1)
    importer.setEpisodeName('Общая')
    assert not importer.canImport
    before = deepcopy(bridge._session.data)
    assert not importer.importSelected()
    assert bridge._session.data == before
    importer.setSelected(1, False)
    assert importer.canImport
    assert importer.importSelected()
    assert list(bridge._session.data['episodes']) == ['Общая']


def test_source_regeneration_uses_exact_sheet_mapping_and_first_row(bridge, tmp_path):
    path = make_workbook(tmp_path / 'episodes.xlsx')
    importer = bridge.spreadsheetImport
    assert importer.load(str(path))
    importer.setFirstRow(4)
    assert importer.importSelected()
    second = '02 · Продолжение'
    assert bridge._project_files._parse_episode_source(str(path), bridge._session.data, second)[0]['text'] == 'Другой лист'
    bridge._session.current_episode = second
    bridge.projectFiles.regenerateWorkingText(second)
    assert bridge._script_text_service.load_episode_lines(bridge._session.data, second)[0]['text'] == 'Другой лист'
    first = '01 · Начало'
    assert len(bridge._project_files._parse_episode_source(str(path), bridge._session.data, first)) == 1
    workbook = __import__('openpyxl').load_workbook(path)
    workbook.remove(workbook[second])
    workbook.save(path)
    with pytest.raises(ValueError, match='больше нет листа'):
        bridge._project_files._parse_episode_source(str(path), bridge._session.data, second)


def test_invalid_timing_blocks_entire_batch(bridge, tmp_path):
    path = make_workbook(tmp_path / 'episodes.xlsx')
    workbook = __import__('openpyxl').load_workbook(path)
    workbook['02 · Продолжение']['C2'] = 'bad time'
    workbook.save(path)
    assert bridge.spreadsheetImport.load(str(path))
    before = deepcopy(bridge._session.data)
    assert not bridge.spreadsheetImport.canImport
    assert not bridge.spreadsheetImport.importSelected()
    assert bridge._session.data == before
    bridge.spreadsheetImport.setSelected(1, False)
    assert bridge.spreadsheetImport.importSelected()


def test_csv_options_reload_and_failed_load_keep_configuration(bridge, tmp_path):
    path = tmp_path / 'Episode_05.csv'
    path.write_text('Персонаж;Текст\nГерой;Привет\n', encoding='cp1251')
    importer = bridge.spreadsheetImport
    assert importer.load(str(path))
    importer.setEpisodeName('Моя серия')
    assert importer.setCsvOptions('cp1251', ';')
    assert importer.episodeName == 'Моя серия'
    assert not importer.setCsvOptions('utf-8', ',')
    assert importer.canImport
    assert importer.mapping['text'] == 1
    assert importer.importSelected()
    source = bridge._session.data['script_storage']['episodes']['Моя серия']['source']
    assert source['import_config']['encoding'] == 'cp1251'
    assert source['import_config']['delimiter'] == ';'


def test_import_rejects_audiobooks_and_pending_invalid_edits(bridge, tmp_path, monkeypatch):
    path = make_workbook(tmp_path / 'episodes.xlsx')
    importer = bridge.spreadsheetImport
    assert importer.load(str(path))
    monkeypatch.setattr(bridge._session, 'flush_pending_edits', lambda: False)
    assert not importer.importSelected()
    assert not bridge._session.data['episodes']
    bridge._session.data['project_kind'] = 'audiobook'
    assert not importer.load(str(path))
    assert not importer.importSelected()


def test_docx_raw_preview_first_row_and_saved_table_selection(bridge, tmp_path):
    from docx import Document
    doc = Document()
    for text in ('Первая таблица', 'Вторая таблица'):
        table = doc.add_table(rows=3, cols=2)
        for r, values in enumerate([['Персонаж', 'Текст'], ['Герой', 'Пропустить'], ['Герой', text]]):
            for c, value in enumerate(values):
                table.cell(r,c).text = value
    path = tmp_path / 'Episode_06.docx'
    doc.save(path)
    importer = bridge.docxImport
    assert importer.load(str(path))
    importer.setTable(1)
    importer.setMapping('time_split', -1)
    importer.setFirstRow(3)
    assert len(importer.rawRows) == 3
    assert importer.previewModel.rowCount() == 1
    assert importer.importEpisode('6', False)
    lines = bridge._project_files._parse_episode_source(str(path), bridge._session.data, '6')
    assert [row['text'] for row in lines] == ['Вторая таблица']


def test_frame_timing_uses_project_fps_without_mutating_preview(bridge, tmp_path):
    from services.project_fps_service import set_project_fps
    set_project_fps(bridge._session.data, 24)
    path = tmp_path / 'frames.csv'
    path.write_text('Персонаж;Начало;Конец;Текст\nГерой;00:00:02:12;00:00:03:12;Кадры\n', encoding='utf-8')
    before = deepcopy(bridge._session.data)
    importer = bridge.spreadsheetImport
    assert importer.load(str(path))
    assert bridge._session.data == before
    assert importer.fps == 24
    importer.autoDetect()
    assert importer.fps == 24
    assert importer.importSelected()
    line = bridge._script_text_service.load_episode_lines(bridge._session.data, '1')[0]
    assert line['s'] == 2.5
    assert line['e'] == 3.5


def test_header_checkbox_updates_preview_and_saved_csv_source(bridge, tmp_path):
    path = tmp_path / 'Episode_07.csv'
    path.write_text('sep=;\nКомментарий\nПерсонаж;Текст\nГерой;Реплика\n', encoding='utf-8')
    importer = bridge.spreadsheetImport
    assert importer.load(str(path))
    assert importer.ignoreHeader
    assert importer.tableStartRow == 2
    assert importer.firstRow == 3
    importer.setIgnoreHeader(False)
    assert importer.tableStartRow == 2
    assert importer.firstRow == 2
    assert importer.previewModel.rows()[0]['text'] == 'Текст'
    importer.setIgnoreHeader(True)
    assert importer.firstRow == 3
    assert [row['text'] for row in importer.previewModel.rows()] == ['Реплика']
    assert importer.importSelected()
    source = bridge._session.data['script_storage']['episodes']['07']['source']
    assert source['import_config']['ignore_header'] is True
    assert source['import_config']['first_row'] == 3
    assert [row['text'] for row in bridge._project_files._parse_episode_source(str(path), bridge._session.data, '07')] == ['Реплика']


def test_header_checkbox_without_detected_header_and_per_sheet_state(bridge, tmp_path):
    path = tmp_path / 'episodes.xlsx'
    workbook = Workbook()
    first = workbook.active
    first.title = '1'
    first.append(['Person', 'Dialogue'])
    first.append(['Hero', 'First'])
    second = workbook.create_sheet('2')
    second.append(['Guest', 'Second'])
    second.append(['Guest', 'Third'])
    workbook.save(path)
    importer = bridge.spreadsheetImport
    assert importer.load(str(path))
    importer.setIgnoreHeader(True)
    assert importer.firstRow == 2
    assert importer.tableStartRow == 1
    importer.setSheet(1)
    assert importer.ignoreHeader is False
    assert importer.firstRow == 1
    importer.setSheet(0)
    assert importer.ignoreHeader is True
    assert importer.firstRow == 2
    importer.setTableStartRow(2)
    assert importer.firstRow == 3
    importer.setIgnoreHeader(False)
    assert importer.firstRow == 2
    assert importer.canImport


@pytest.mark.parametrize('ignore_header', [False, True])
def test_apply_columns_to_all_copies_header_flag_and_keeps_start_rows(bridge, tmp_path, ignore_header):
    path = make_workbook(tmp_path / 'episodes.xlsx')
    importer = bridge.spreadsheetImport
    assert importer.load(str(path))
    importer.setEpisodeName('Custom')
    importer.setMapping('time_start', -1)
    importer.setMapping('time_end', -1)
    importer.setMapping('text', 1)
    importer.setMapping('character', 0)
    importer.setIgnoreHeader(ignore_header)
    importer.setTimeUnit('excel')
    importer.setFps(24)
    importer.setSheet(1)
    importer.setIgnoreHeader(not ignore_header)
    importer.setSheet(0)
    originals = deepcopy(importer._states)
    assert importer.canApplyMappingToAll
    importer.applyMappingToAll()
    for i, state in enumerate(importer._states):
        assert state['episode'] == originals[i]['episode']
        assert state['selected'] == originals[i]['selected']
        for key in ('time_unit', 'fps'):
            assert state['config'][key] == originals[i]['config'][key]
        if importer._sheets[i].rows:
            assert state['config']['mapping'] == importer.mapping
            original_config = originals[i]['config']
            start = original_config['first_row'] - int(original_config['ignore_header'])
            assert state['config']['ignore_header'] is ignore_header
            assert state['config']['first_row'] == start + int(ignore_header)
        else:
            assert state['config'] == originals[i]['config']
    importer.setSheet(1)
    assert importer.mapping['text'] == 1
    importer.setMapping('text', 0)
    importer.setSheet(0)
    assert importer.mapping['text'] == 1  # Copies, not shared dictionaries.
    assert not bridge._session.data['episodes']


def test_apply_columns_flags_incompatible_selected_sheets(bridge, tmp_path):
    path = make_workbook(tmp_path / 'episodes.xlsx')
    importer = bridge.spreadsheetImport
    assert importer.load(str(path))
    before = deepcopy(bridge._session.data)
    importer.applyMappingToAll()
    assert not importer.canImport
    assert '02 · Продолжение' in importer.validationMessage
    importer.setSheet(1)
    assert importer.mapping['text'] == 3
    assert not importer.importSelected()
    assert bridge._session.data == before
    importer.setSelected(1, False)
    assert importer.canImport
    assert importer.importSelected()


def test_docx_header_checkbox_is_per_table_and_survives_regeneration(bridge, tmp_path):
    from docx import Document
    doc = Document()
    for value in ('Первый', 'Второй'):
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Персонаж'
        table.cell(0, 1).text = 'Текст'
        table.cell(1, 0).text = 'Герой'
        table.cell(1, 1).text = value
    path = tmp_path / 'Episode_08.docx'
    doc.save(path)
    importer = bridge.docxImport
    assert importer.load(str(path))
    assert importer.ignoreHeader
    assert importer.tableStartRow == 1
    importer.setIgnoreHeader(False)
    assert importer.firstRow == 1
    assert importer.previewModel.rows()[0]['text'] == 'Текст'
    importer.setTable(1)
    assert importer.ignoreHeader
    assert importer.firstRow == 2
    importer.setTable(0)
    assert not importer.ignoreHeader
    assert importer.firstRow == 1
    importer.setMapping('time_split', -1)
    assert importer.importEpisode('8', True)
    texts = [row['text'] for row in bridge._project_files._parse_episode_source(str(path), bridge._session.data, '8')]
    assert texts == ['Текст', 'Первый', 'Второй']
