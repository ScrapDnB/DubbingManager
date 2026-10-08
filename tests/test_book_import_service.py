from services.book_import_service import BookImportService


def test_split_texts_into_chapters_by_headings():
    service = BookImportService()

    chapters = service.split_texts_into_chapters([
        {
            "page": 1,
            "text": "Глава 1\nПервый абзац.\n\nВторой абзац.",
        },
        {
            "page": 2,
            "text": "Глава 2\nТретий абзац.",
        },
    ])

    assert [chapter.title for chapter in chapters] == ["Глава 1", "Глава 2"]
    assert chapters[0].paragraphs == ["Первый абзац.", "Второй абзац."]
    assert chapters[1].paragraphs == ["Третий абзац."]


def test_split_texts_without_headings_creates_single_chapter():
    service = BookImportService()

    chapters = service.split_texts_into_chapters([
        {"page": 1, "text": "Первый абзац.\n\nВторой абзац."},
    ])

    assert len(chapters) == 1
    assert chapters[0].title == "Глава 1"
    assert chapters[0].paragraphs == ["Первый абзац.", "Второй абзац."]


def test_named_chapters_keep_their_titles_and_do_not_shift_numbering():
    service = BookImportService()

    chapters = service.split_texts_into_chapters([
        {
            "page": 1,
            "text": (
                "Вступление\nПеред началом.\n\n"
                "Пролог\nНачало истории.\n\n"
                "Глава 1\nПервая глава."
            ),
        },
    ])

    assert [chapter.title for chapter in chapters] == [
        "Вступление",
        "Пролог",
        "Глава 1",
    ]
    assert chapters[0].paragraphs == ["Перед началом."]
    assert chapters[1].paragraphs == ["Начало истории."]
    assert chapters[2].paragraphs == ["Первая глава."]


def test_table_of_contents_is_not_imported_as_chapter():
    service = BookImportService()

    chapters = service.split_texts_into_chapters([
        {
            "page": 1,
            "text": (
                "Оглавление\n"
                "Пролог 5\n"
                "Глава 1 12\n"
                "Глава 2 20\n\n"
                "Пролог\n"
                "Настоящий текст пролога.\n\n"
                "Глава 1\n"
                "Настоящий текст первой главы."
            ),
        },
    ])

    assert [chapter.title for chapter in chapters] == ["Пролог", "Глава 1"]
    assert chapters[0].paragraphs == ["Настоящий текст пролога."]
    assert chapters[1].paragraphs == ["Настоящий текст первой главы."]


def test_trailing_table_of_contents_entries_are_not_imported_as_chapters():
    service = BookImportService()

    chapters = service.split_texts_into_chapters([
        {
            "page": 1,
            "text": (
                "Вступление\n"
                "Текст вступления.\n\n"
                "Глава 1\n"
                "Текст первой главы.\n\n"
                "Оглавление\n"
                "Вступление 5\n"
                "Пролог 8\n"
                "Глава 1 12\n"
                "Глава 2 ........ 20"
            ),
        },
    ])

    assert [chapter.title for chapter in chapters] == ["Вступление", "Глава 1"]
    assert chapters[0].paragraphs == ["Текст вступления."]
    assert chapters[1].paragraphs == ["Текст первой главы."]


def test_build_lines_from_segments_defaults_to_narrator():
    service = BookImportService()

    lines = service.build_lines_from_segments([
        {"character": "Автор", "text": "Авторский текст."},
        {"character": "Иван", "text": "Привет."},
        {"text": "Снова автор."},
    ])

    assert [line["char"] for line in lines] == ["Автор", "Иван", "Автор"]
    assert [line["id"] for line in lines] == [0, 1, 2]
    assert all(line["_book_text"] is True for line in lines)


def test_pdf_visual_lines_are_reconstructed_as_book_paragraphs():
    service = BookImportService()

    paragraphs = service._paragraphs_from_text(
        "151\n"
        "ПАДАЯ, СЛОВНО ЗВЕЗДЫ\n"
        "ЗАКАРИ\n"
        "— Я не хотел подслушивать, но у этой\n"
        "девушки было много предположений о нас\n"
        "с тобой.\n"
        "— Она не умеет не лезть не в свое дело.\n"
        "Он поднимает на меня взгляд, теплый и вни-\n"
        "мательный."
    )

    assert paragraphs == [
        "ЗАКАРИ",
        "— Я не хотел подслушивать, но у этой девушки было много "
        "предположений о нас с тобой.",
        "— Она не умеет не лезть не в свое дело.",
        "Он поднимает на меня взгляд, теплый и внимательный.",
    ]


def test_paragraph_continuing_on_next_page_is_joined():
    service = BookImportService()

    chapters = service.split_texts_into_chapters([
        {"page": 1, "text": "Глава 1\nБелое небо такое же, как"},
        {"page": 2, "text": "2\nНАЗВАНИЕ КНИГИ\nи земля."},
    ])

    assert chapters[0].paragraphs == ["Белое небо такое же, как и земля."]


def test_custom_chapter_keywords_replace_defaults():
    service = BookImportService(["Часть"])

    chapters = service.split_texts_into_chapters([
        {
            "page": 1,
            "text": (
                "Глава 1\nНе должна стать отдельной главой.\n\n"
                "Часть 2\nДолжна стать новой главой."
            ),
        },
    ])

    assert [chapter.title for chapter in chapters] == ["Глава 1", "Часть 2"]
    assert chapters[0].paragraphs == [
        "Глава 1 Не должна стать отдельной главой.",
    ]


def test_custom_keyword_without_number_does_not_consume_next_line():
    service = BookImportService(["Плейлист"])

    chapters = service.split_texts_into_chapters([
        {
            "page": 1,
            "text": (
                "Плейлист\n"
                "Первый трек.\n"
                "Второй трек."
            ),
        },
    ])

    assert [chapter.title for chapter in chapters] == ["Плейлист"]
    assert chapters[0].paragraphs == ["Первый трек.", "Второй трек."]


def test_empty_chapter_keywords_disable_automatic_splitting():
    service = BookImportService([])

    chapters = service.split_texts_into_chapters([
        {"page": 1, "text": "Глава 1\nТекст.\n\nГлава 2\nПродолжение."},
    ])

    assert len(chapters) == 1


def test_chapters_to_html_keeps_full_imported_book_text():
    service = BookImportService()
    chapters = service.split_texts_into_chapters([
        {"page": 1, "text": "Глава 1\nПервый текст."},
        {"page": 2, "text": "Глава 2\nВторой текст."},
    ])

    html_text = service.chapters_to_html(chapters)

    assert "Глава 1" in html_text
    assert "Первый текст." in html_text
    assert "Глава 2" in html_text
    assert "Второй текст." in html_text


def test_import_utf8_and_windows1251_prose_txt(tmp_path):
    service = BookImportService()
    for encoding in ("utf-8-sig", "cp1251"):
        path = tmp_path / "book.txt"
        path.write_bytes("Пролог\nНачало книги.\n\nГлава 1\nПервый абзац.\n\nВторой абзац.".encode(encoding))
        chapters = service.import_file(str(path))
        assert [c.title for c in chapters] == ["Пролог", "Глава 1"]
        assert chapters[1].paragraphs == ["Первый абзац.", "Второй абзац."]


def test_import_docx_prose_uses_headings_without_screenplay_columns(tmp_path):
    from docx import Document
    document = Document()
    document.add_heading("Обсерватория", level=1)
    document.add_paragraph("Первый абзац книги.")
    document.add_paragraph("— Кто здесь?")
    document.add_heading("Возвращение", level=1)
    document.add_paragraph("Последний абзац.")
    path = tmp_path / "book.docx"
    document.save(path)
    chapters = BookImportService().import_file(str(path))
    assert [c.title for c in chapters] == ["Обсерватория", "Возвращение"]
    assert chapters[0].paragraphs == ["Первый абзац книги.", "— Кто здесь?"]


def test_import_epub_follows_spine_and_skips_navigation(tmp_path):
    from zipfile import ZipFile
    path = tmp_path / "book.epub"
    with ZipFile(path, "w") as archive:
        archive.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="OPS/content.opf"/></rootfiles></container>')
        archive.writestr("OPS/content.opf", '<package><manifest><item id="second" href="second.xhtml"/><item id="first" href="first.xhtml"/><item id="nav" href="nav.xhtml" properties="nav"/></manifest><spine><itemref idref="nav"/><itemref idref="first"/><itemref idref="second"/></spine></package>')
        archive.writestr("OPS/second.xhtml", '<html><body><h1>Глава 2</h1><p>Второй текст.</p></body></html>')
        archive.writestr("OPS/first.xhtml", '<html><body><h1>Глава 1</h1><p>Первый текст.</p></body></html>')
        archive.writestr("OPS/nav.xhtml", '<html><body><p>Оглавление.</p></body></html>')
    chapters = BookImportService().import_file(str(path))
    assert [c.title for c in chapters] == ["Глава 1", "Глава 2"]
    assert chapters[0].paragraphs == ["Первый текст."]
